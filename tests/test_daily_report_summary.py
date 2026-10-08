from __future__ import annotations

import csv
import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PyPDF2 import PdfReader
from PyPDF2.generic import ContentStream

from dashboard.api import DAILY_REPORT_ENDPOINT, route_dashboard_request
from dashboard.daily_report_docx import build_summary_docx_report
from dashboard.daily_report_pdf import build_summary_pdf_report
from dashboard.daily_report_summary import summary_sections
from dashboard.expanded_daily_report import _read_comparison, build_daily_report_model
from dashboard.expanded_signal_board import build_expanded_signal_board
from src.expanded_shadow_ledger import LEDGER_FIELDS
from src.expanded_shadow_ops import ExpandedShadowPaths

ROOT = Path(__file__).resolve().parents[1]
DAY = "2026-10-08"
PRIOR = "2026-10-07"


def _write_csv(path: Path, fields, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _signal(day, ticker="000001", **overrides):
    row = {field: "" for field in LEDGER_FIELDS}
    row.update(
        basDd=day, signal_date=day, stock_code=ticker, stock_name=f"종목{ticker}",
        signal_price="10000", signal_score="80", signal_type="BUY_WATCH",
        decision="CANDIDATE", foreign_status="POSITIVE", foreign_5d_ratio="0.2",
    )
    row.update(overrides)
    return row


def _comparison_root(tmp_path, prior_rows, today_rows, *, previous=PRIOR):
    paths = ExpandedShadowPaths(tmp_path)
    _write_csv(paths.market_dir(DAY) / "005930.csv", ["date"], [{"date": previous}, {"date": DAY}])
    _write_csv(paths.signal_ledger_path, LEDGER_FIELDS, prior_rows + today_rows)
    for day, rows in ((previous, prior_rows), (DAY, today_rows)):
        path = paths.manifest_path(day)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps({
            "basDd": day, "status": "SUCCESS", "ready_count": 574,
            "canonical_universe_count": 574, "signal_count": len(rows),
            "new_candidate_count": sum(row["decision"] == "CANDIDATE" for row in rows),
            "attempted_ticker_count": 574, "run_id": day,
        }), encoding="utf-8")
    return paths


def _pdf_text(binary):
    reader = PdfReader(BytesIO(binary))
    parts = []
    for page in reader.pages:
        for operands, operator in ContentStream(page.get_contents(), reader).operations:
            if operator not in (b"Tj", b"TJ"):
                continue
            values = operands if operator == b"Tj" else operands[0]
            for value in values:
                if hasattr(value, "original_bytes"):
                    parts.append(value.original_bytes.decode("utf-16-be"))
    return "".join(parts)


def _squash(text):
    return "".join(text.split())


def test_real_summary_has_every_candidate_and_every_overheated_signal():
    model = build_daily_report_model(ROOT, DAY)
    sections = summary_sections(model)
    assert len(sections[0].rows) == 18
    assert {row[0] for row in sections[0].rows} == {
        f"{row.stock_name} ({row.ticker})" for row in model.new_candidates
    }
    assert len(sections[2].rows) == 7
    assert sum(row[2] == "EXCLUDED" for row in sections[2].rows) == 3
    assert any(row[0] == "LG에너지솔루션 (373220)" and row[2:5] == ("403,000", "89.2", "STRONG_WATCH") for row in sections[0].rows)
    assert model.run_summary.source_date == DAY
    assert model.performance_as_of == DAY
    assert sections[3].rows[0][1] == "123"
    assert sections[3].rows[1][1] == "50"
    assert sections[3].rows[2][1:3] == ("0", "확인 불가")
    board = build_expanded_signal_board(ROOT)
    assert board["run_summary"]["source_date"] == model.run_summary.source_date
    assert {
        (row.stock_name, row.ticker, row.market, row.signal_date, row.entry_price, row.signal_score, row.foreign_status)
        for row in model.new_candidates
    } == {
        (row["stock_name"], row["ticker"], row["market"], row["signal_date"], row["entry_price"], row["signal_score"], row["foreign_status"])
        for row in board["new_candidates"]["records"]
    }


def test_real_previous_day_is_event_only_comparison():
    comparison = build_daily_report_model(ROOT, DAY).comparison
    assert comparison.status == "AVAILABLE"
    assert comparison.previous_date == PRIOR
    assert len(comparison.changes) == 27
    assert all(row.change == "신규" for row in comparison.changes)
    assert len(comparison.previous_only) == 17
    assert "상태 지속은 추정하지" in comparison.reason


@pytest.mark.parametrize("field,value", [
    ("signal_score", "81"), ("signal_type", "OVERHEATED"),
    ("decision", "EXCLUDED"), ("foreign_status", "NEGATIVE"),
    ("foreign_5d_ratio", "0.3"),
])
def test_changed_signal_uses_stored_comparison_fields(tmp_path, field, value):
    paths = _comparison_root(tmp_path, [_signal(PRIOR)], [_signal(DAY, **{field: value})])
    comparison = _read_comparison(paths, DAY)
    assert comparison.status == "AVAILABLE"
    assert comparison.changes[0].change == "변경"
    assert value in comparison.changes[0].current


def test_maintained_compares_numeric_values_not_string_formatting(tmp_path):
    paths = _comparison_root(tmp_path, [_signal(PRIOR)], [_signal(DAY, signal_score="80.0", foreign_5d_ratio="0.20")])
    assert _read_comparison(paths, DAY).changes[0].change == "유지"


def test_new_and_previous_only_are_not_inferred_states(tmp_path):
    paths = _comparison_root(tmp_path, [_signal(PRIOR)], [_signal(DAY, "000002")])
    comparison = _read_comparison(paths, DAY)
    assert [row.change for row in comparison.changes] == ["신규"]
    assert comparison.previous_only == ("종목000001(000001)",)


@pytest.mark.parametrize("bad", ["missing_manifest", "failed", "incomplete", "count", "duplicate", "missing_calendar", "invalid_pointer", "invalid_manifest"])
def test_unavailable_comparison_never_falls_back_or_estimates(tmp_path, bad):
    paths = _comparison_root(tmp_path, [_signal(PRIOR)], [_signal(DAY)])
    manifest = paths.manifest_path(PRIOR)
    if bad == "missing_manifest":
        manifest.unlink()
    elif bad == "missing_calendar":
        (paths.market_dir(DAY) / "005930.csv").unlink()
    elif bad == "duplicate":
        _write_csv(paths.signal_ledger_path, LEDGER_FIELDS, [_signal(PRIOR), _signal(DAY), _signal(DAY)])
        current = json.loads(paths.manifest_path(DAY).read_text())
        current["signal_count"] = 2
        paths.manifest_path(DAY).write_text(json.dumps(current))
    elif bad == "invalid_pointer":
        pointer = paths.latest_manifest_path(PRIOR)
        pointer.parent.mkdir(parents=True, exist_ok=True)
        pointer.write_text("{}")
    elif bad == "invalid_manifest":
        manifest.write_text("[]")
    else:
        payload = json.loads(manifest.read_text())
        payload.update({"status": "FAILED"} if bad == "failed" else {"ready_count": 573} if bad == "incomplete" else {"signal_count": 2})
        manifest.write_text(json.dumps(payload))
    comparison = _read_comparison(paths, DAY)
    assert comparison.status == "UNAVAILABLE"
    assert comparison.reason.startswith("비교 불가:")
    assert comparison.changes == ()


def test_zero_signal_previous_day_is_available(tmp_path):
    paths = _comparison_root(tmp_path, [], [_signal(DAY)])
    comparison = _read_comparison(paths, DAY)
    assert comparison.status == "AVAILABLE"
    assert comparison.changes[0].change == "신규"


def test_zero_signal_today_can_generate_summary(tmp_path):
    paths = _comparison_root(tmp_path, [], [])
    model = build_daily_report_model(paths.repo_root, DAY)
    assert model.new_candidates_status == "EMPTY"
    status, _, _ = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={DAY}&format=docx", tmp_path)
    assert status == 200
    assert model.comparison.status == "AVAILABLE"


def test_zero_signal_manifest_does_not_mask_missing_ledger(tmp_path):
    paths = _comparison_root(tmp_path, [], [])
    paths.signal_ledger_path.unlink()
    status, _, _ = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={DAY}&format=docx", tmp_path)
    assert status == 404


def test_previous_day_comes_from_stored_trading_dates_not_calendar_subtraction(tmp_path):
    friday = "2026-10-02"
    paths = _comparison_root(tmp_path, [_signal(friday)], [_signal(DAY)], previous=friday)
    assert _read_comparison(paths, DAY).previous_date == friday


def test_missing_immediate_previous_run_does_not_use_older_run(tmp_path):
    paths = _comparison_root(tmp_path, [_signal(PRIOR)], [_signal(DAY)])
    paths.manifest_path(PRIOR).unlink()
    paths.manifest_path("2026-10-06").write_text(json.dumps({"basDd": "2026-10-06", "status": "SUCCESS"}))
    assert _read_comparison(paths, DAY).status == "UNAVAILABLE"


def test_historical_summary_uses_historical_run_and_never_future_returns():
    model = build_daily_report_model(ROOT, "2026-09-18")
    assert model.run_summary.source_date == "2026-09-18"
    performance = summary_sections(model)[3]
    assert all(row[1:] == ("확인 불가",) * 4 for row in performance.rows)
    assert "이후 데이터 대체 금지" in performance.note


def test_summary_docx_pdf_content_is_exactly_shared_and_pdf_is_two_pages():
    model = build_daily_report_model(ROOT, DAY)
    docx = Document(BytesIO(build_summary_docx_report(model)))
    pdf = build_summary_pdf_report(model)
    assert len(PdfReader(BytesIO(pdf)).pages) == 2
    docx_text = "\n".join(
        [paragraph.text for paragraph in docx.paragraphs]
        + [cell.text for table in docx.tables for row in table.rows for cell in row.cells]
    )
    pdf_text = _pdf_text(pdf)
    for section in summary_sections(model):
        for text in (section.title, section.note, *section.headers, *(value for row in section.rows for value in row)):
            assert _squash(text) in _squash(docx_text)
            assert _squash(text) in _squash(pdf_text)
    assert docx.styles["Normal"].font.name == "Malgun Gothic"


def test_return_summary_ignores_unobserved_values_not_as_zero():
    model = build_daily_report_model(ROOT, DAY)
    record = next(row for row in model.performance if row.return_5d is not None)
    partial = replace(record, return_5d=0.0, excess_5d=None)
    missing = replace(record, ticker="999999", return_5d=None, excess_5d=99)
    performance = summary_sections(replace(model, performance=[partial, missing]))[3]
    assert performance.rows[0] == ("5D", "1", "+0.00%", "0", "확인 불가")


def test_summary_keeps_fractional_signal_price_without_rounding():
    model = build_daily_report_model(ROOT, DAY)
    candidate = replace(model.new_candidates[0], entry_price=12345.75)
    assert summary_sections(replace(model, new_candidates=[candidate]))[0].rows[0][2] == "12,345.75"


@pytest.mark.parametrize("mode", ["", "invalid", "summary&mode=detail"])
def test_api_rejects_invalid_or_ambiguous_mode(mode):
    status, _, body = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={DAY}&format=pdf&mode={mode}", ROOT)
    assert status == 400
    assert json.loads(body)["error_code"] == "INVALID_MODE"


@pytest.mark.parametrize("format", ["docx", "pdf"])
def test_api_summary_is_default_and_detail_remains_selectable(format):
    default_status, default_headers, default = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={DAY}&format={format}", ROOT)
    explicit_status, _, explicit = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={DAY}&format={format}&mode=summary", ROOT)
    detail_status, detail_headers, detail = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={DAY}&format={format}&mode=detail", ROOT)
    assert default_status == explicit_status == detail_status == 200
    assert "_summary." in default_headers["Content-Disposition"]
    assert "_summary." not in detail_headers["Content-Disposition"]
    text = lambda body: _pdf_text(body) if format == "pdf" else "\n".join(p.text for p in Document(BytesIO(body)).paragraphs)
    assert _squash(text(default)) == _squash(text(explicit))
    assert "쉬운 종목 분석" in text(detail)
    assert len(detail) > len(default)

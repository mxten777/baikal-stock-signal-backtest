from __future__ import annotations

import hashlib
import subprocess
from io import BytesIO
from pathlib import Path

from docx import Document
from PyPDF2 import PdfReader

from dashboard.api import DAILY_REPORT_ENDPOINT, route_dashboard_request

REPO_ROOT = Path(__file__).resolve().parents[1]
REAL_SOURCE_DATE = "2026-09-18"


def _pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(pdf_bytes))
    raw = "".join(page.extract_text() for page in reader.pages).replace("\x00", "")
    return "".join(raw.split())


def _docx_all_text(docx_bytes: bytes) -> str:
    document = Document(BytesIO(docx_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# 1: DOCX endpoint 200 + correct MIME
def test_docx_endpoint_returns_200_with_correct_mime():
    status, headers, body = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=docx", REPO_ROOT
    )

    assert status == 200
    assert headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert len(body) > 0
    Document(BytesIO(body))  # must parse as a valid docx


# 2: PDF endpoint 200 + correct MIME
def test_pdf_endpoint_returns_200_with_correct_mime():
    status, headers, body = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=pdf", REPO_ROOT
    )

    assert status == 200
    assert headers["Content-Type"] == "application/pdf"
    assert len(body) > 0
    reader = PdfReader(BytesIO(body))
    assert len(reader.pages) >= 1


# 3: Content-Disposition filename/date
def test_content_disposition_filename_matches_date_and_format():
    _, docx_headers, _ = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=docx", REPO_ROOT
    )
    _, pdf_headers, _ = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=pdf", REPO_ROOT
    )

    assert docx_headers["Content-Disposition"] == f'attachment; filename="BAIKAL_Daily_Report_{REAL_SOURCE_DATE}.docx"'
    assert pdf_headers["Content-Disposition"] == f'attachment; filename="BAIKAL_Daily_Report_{REAL_SOURCE_DATE}.pdf"'


# 4: real 2026-09-18 report generation (both formats)
def test_real_2026_09_18_report_generation():
    docx_status, _, docx_body = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=docx", REPO_ROOT
    )
    pdf_status, _, pdf_body = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=pdf", REPO_ROOT
    )

    assert docx_status == 200 and pdf_status == 200
    assert REAL_SOURCE_DATE in _docx_all_text(docx_body)
    assert REAL_SOURCE_DATE in _pdf_text(pdf_body)


# 5: docx has 23 candidates + signal score / entry price / foreign status
def test_docx_reflects_23_candidates_with_required_fields():
    status, _, body = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=docx", REPO_ROOT
    )
    assert status == 200

    document = Document(BytesIO(body))
    candidate_tables = [t for t in document.tables if len(t.rows[0].cells) == 7]
    assert len(candidate_tables) == 1
    table = candidate_tables[0]
    assert len(table.rows) == 24  # header + 23 candidates

    header_cells = [cell.text for cell in table.rows[0].cells]
    assert "Signal Score" in header_cells
    assert "Entry Price" in header_cells
    assert "Foreign Status" in header_cells
    for row in table.rows[1:]:
        values = [cell.text for cell in row.cells]
        assert values[4] != ""  # Entry Price
        assert values[5] != ""  # Signal Score
        assert values[6] in ("POSITIVE", "NEGATIVE", "NEUTRAL")  # Foreign Status


# 6: PDF generation confirmed (already covered by test 2, kept explicit per spec item 6)
def test_pdf_generation_confirmed_for_real_date():
    status, headers, body = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=pdf", REPO_ROOT
    )

    assert status == 200
    assert headers["Content-Type"] == "application/pdf"
    reader = PdfReader(BytesIO(body))
    assert len(reader.pages) >= 1


# 7: missing/invalid date -> controlled error
def test_missing_date_returns_400():
    status, headers, _ = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?format=docx", REPO_ROOT)

    assert status == 400
    assert headers["Content-Type"].startswith("application/json")


def test_invalid_date_format_returns_400():
    status, _, _ = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date=2026/09/18&format=docx", REPO_ROOT)
    assert status == 400


def test_invalid_calendar_date_returns_400():
    status, _, _ = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date=2026-13-40&format=docx", REPO_ROOT)
    assert status == 400


# 8: invalid format -> 400
def test_missing_format_returns_400():
    status, _, _ = route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}", REPO_ROOT)
    assert status == 400


def test_invalid_format_value_returns_400():
    status, _, _ = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=xlsx", REPO_ROOT
    )
    assert status == 400


# 9: no fallback to another date - a date with no ledger data must 404, never silently substitute
def test_unavailable_date_returns_404_without_fallback():
    status, headers, _ = route_dashboard_request(
        "GET", f"{DAILY_REPORT_ENDPOINT}?date=2099-01-01&format=docx", REPO_ROOT
    )

    assert status == 404
    assert headers["Content-Type"].startswith("application/json")


# 10: protected artifacts unchanged after API calls
def _sha256(path: Path) -> str | None:
    if not path.is_file():
        return None
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_protected_artifacts_unchanged_after_api_calls():
    protected = [
        REPO_ROOT / "output/expanded_shadow/expanded_shadow_run.json",
        REPO_ROOT / "output/expanded_shadow/expanded_shadow_signal_ledger.csv",
        REPO_ROOT / "output/expanded_shadow/expanded_candidate_performance_ledger.csv",
    ]
    before = {path: _sha256(path) for path in protected}

    route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=docx", REPO_ROOT)
    route_dashboard_request("GET", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=pdf", REPO_ROOT)

    after = {path: _sha256(path) for path in protected}
    assert before == after

    protected_paths = [
        "src",
        "scripts",
        "dashboard/operations.py",
        "dashboard/daily_signal_board.py",
        "dashboard/expanded_signal_board.py",
        "dashboard/dual_shadow.py",
        "dashboard/adapter",
        "dashboard/runner",
        "dashboard/daily_report_model.py",
        "dashboard/expanded_daily_report.py",
        "dashboard/daily_report_docx.py",
        "dashboard/daily_report_pdf.py",
        ":(exclude)scripts/daily_scheduler.py",
        ":(exclude)scripts/safe_investor_update.py",
        ":(exclude)scripts/daily_operational_run.py",
        ":(exclude)scripts/daily_health_report.py",
        ":(exclude)src/expanded_shadow_pipeline.py",
    ]
    completed = subprocess.run(
        ["git", "diff", "--name-only", "--", *protected_paths],
        check=True,
        capture_output=True,
        text=True,
        cwd=REPO_ROOT,
    )
    assert completed.stdout.strip() == ""


# 11: existing dashboard/report regression not broken by routing addition
def test_expanded_signal_board_endpoint_still_works():
    status, headers, body = route_dashboard_request("GET", "/api/dashboard/expanded-shadow", REPO_ROOT)

    assert status == 200
    assert headers["Content-Type"].startswith("application/json")
    assert len(body) > 0


def test_non_get_method_still_rejected_for_daily_report():
    status, headers, _ = route_dashboard_request(
        "POST", f"{DAILY_REPORT_ENDPOINT}?date={REAL_SOURCE_DATE}&format=docx", REPO_ROOT
    )

    assert status == 405
    assert headers["Allow"] == "GET"

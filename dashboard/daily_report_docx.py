"""Read-only DOCX renderer for DailyReportModel.

Takes only a DailyReportModel as input. Never reads artifacts, ledgers, or
manifests directly, and never touches Signal Engine / Production / Shadow /
DUAL / Scheduler / Expanded operational modules.
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from dashboard.daily_report_model import DailyReportModel

try:
    from zoneinfo import ZoneInfo

    _KST = ZoneInfo("Asia/Seoul")
except Exception:  # pragma: no cover - tzdata missing fallback
    _KST = timezone(timedelta(hours=9))

TITLE_LINES = (
    "BAIKAL Stock Signal",
    "Daily Recommendation Report",
    "Expanded Shadow / Research & Monitoring",
)
FOOTER_LINES = (
    "BAIKAL Stock Signal",
    "Expanded Shadow / Research & Monitoring",
    "본 보고서는 연구 및 모니터링을 위한 참고자료이며 투자판단 및 투자결과에 대한 책임은 투자자 본인에게 있습니다.",
)

NEW_CANDIDATE_HEADERS = ("종목명", "Ticker", "Market", "Signal Date", "Entry Price", "Signal Score", "Foreign Status")
NEW_CANDIDATE_COL_WIDTHS_CM = (3.4, 2.0, 1.8, 2.4, 2.4, 2.2, 2.6)
NEW_CANDIDATE_RIGHT_ALIGN = (4, 5)

PERFORMANCE_HEADERS = (
    "종목명",
    "Ticker",
    "Signal Date",
    "Entry Price",
    "Status",
    "5D Return",
    "5D Benchmark",
    "5D Excess",
    "10D Return",
    "10D Benchmark",
    "10D Excess",
    "20D Return",
    "20D Benchmark",
    "20D Excess",
)
PERFORMANCE_RIGHT_ALIGN = (3, 5, 6, 7, 8, 9, 10, 11, 12, 13)


def build_docx_report(model: DailyReportModel) -> bytes:
    document = Document()

    document.add_heading(TITLE_LINES[0], level=0)
    for line in TITLE_LINES[1:]:
        document.add_heading(line, level=1)

    document.add_heading("1. 기준정보", level=1)
    _add_kv_table(
        document,
        (
            ("기준일", _display(model.run_summary.source_date)),
            ("생성일", _generated_at_kst()),
            ("Run ID", _display(model.run_summary.run_id)),
        ),
    )

    document.add_heading("2. Expanded Run Summary", level=1)
    _add_kv_table(
        document,
        (
            ("Universe", _display(model.run_summary.universe)),
            ("Attempted", _display(model.run_summary.attempted)),
            ("READY", _display(model.run_summary.ready)),
            ("Signals", _display(model.run_summary.signals)),
            ("CANDIDATE", _display(model.run_summary.candidate)),
            ("EXCLUDED", _display(model.run_summary.excluded)),
            ("NO_SIGNAL", _display(model.run_summary.no_signal)),
            ("Failure", _display(model.run_summary.failure)),
        ),
    )

    document.add_heading("3. Today's New Candidates", level=1)
    if model.new_candidates:
        table = document.add_table(rows=1, cols=len(NEW_CANDIDATE_HEADERS))
        table.style = "Table Grid"
        _set_header_row(table, NEW_CANDIDATE_HEADERS)
        for record in model.new_candidates:
            cells = table.add_row().cells
            values = (
                record.stock_name,
                record.ticker,
                record.market,
                record.signal_date,
                _format_price(record.entry_price),
                _format_score(record.signal_score),
                record.foreign_status,
            )
            _set_row(cells, values, NEW_CANDIDATE_RIGHT_ALIGN)
        _set_column_widths(table, NEW_CANDIDATE_COL_WIDTHS_CM)
    else:
        document.add_paragraph("신규 후보 없음")

    document.add_heading("4. Candidate Performance Tracking", level=1)
    if model.performance:
        table = document.add_table(rows=1, cols=len(PERFORMANCE_HEADERS))
        table.style = "Table Grid"
        _set_header_row(table, PERFORMANCE_HEADERS)
        for record in model.performance:
            cells = table.add_row().cells
            values = (
                record.stock_name,
                record.ticker,
                record.signal_date,
                _format_price(record.entry_price),
                record.tracking_status,
                _format_percent(record.return_5d),
                _format_percent(record.benchmark_5d),
                _format_percent(record.excess_5d),
                _format_percent(record.return_10d),
                _format_percent(record.benchmark_10d),
                _format_percent(record.excess_10d),
                _format_percent(record.return_20d),
                _format_percent(record.benchmark_20d),
                _format_percent(record.excess_20d),
            )
            _set_row(cells, values, PERFORMANCE_RIGHT_ALIGN)
    else:
        document.add_paragraph("추적 중인 성과 데이터 없음")

    document.add_paragraph()
    for line in FOOTER_LINES:
        paragraph = document.add_paragraph(line)
        paragraph.runs[0].font.size = Pt(9)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_kv_table(document: Document, rows: tuple[tuple[str, str], ...]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value


def _set_header_row(table, headers: tuple[str, ...]) -> None:
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.font.bold = True


def _set_row(cells, values: tuple, right_align_indices: tuple[int, ...]) -> None:
    for index, (cell, value) in enumerate(zip(cells, values)):
        cell.text = str(value)
        if index in right_align_indices:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _set_column_widths(table, widths_cm: tuple[float, ...]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width_cm in zip(row.cells, widths_cm):
            cell.width = Cm(width_cm)


def _generated_at_kst() -> str:
    now = datetime.now(timezone.utc).astimezone(_KST)
    return now.strftime("%Y-%m-%d %H:%M KST")


def _display(value: object) -> str:
    return "—" if value is None else str(value)


def _format_price(value: object) -> str:
    if value is None:
        return "—"
    number = float(value)
    if number.is_integer():
        return f"{int(number):,}"
    return f"{number:,.2f}"


def _format_score(value: object) -> str:
    if value is None:
        return "—"
    return f"{float(value):.1f}"


def _format_percent(value: object) -> str:
    if value is None:
        return "—"
    number = float(value)
    sign = "+" if number > 0 else ""
    return f"{sign}{number:.2f}%"
